"""Evaluate percentile-sweep models on matching resolution groups of the val set.

Idea:
    imgsz values and group bounds both come from the 'design' sheet written by
    analyze_resolution.py (train max_dim percentiles + geometric-mean bounds).
    Each group is named after its matching imgsz -> table group x imgsz has a
    diagonal where scale ~ 1; below it images are shrunk, above it enlarged.
"""

import os
import pathlib

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

os.environ["YOLO_CONFIG_DIR"] = str(pathlib.Path("~/VINH_TRUONG/yolo26/.ultralytics_config").expanduser())
from ultralytics import YOLO  # import after YOLO_CONFIG_DIR is set

# ===== CONFIG =====
ROOT = pathlib.Path("~/VINH_TRUONG/yolo26").expanduser()
RES_XLSX = ROOT / "resolution_outputs" / "resolution_per_image_train.xlsx"  # sheets: resolution, design
SPLIT_DIR = ROOT / "dataset" / "graz" / "split"
GROUP_DIR = ROOT / "dataset" / "graz" / "percentile_groups"          # per-group .txt + .yaml
SWEEP_DIR   = ROOT / "runs" / "multi_size_first_test"
EVAL_DIR    = ROOT / "runs" / "eval" / "multi_size_first_test" / "percentile_5_groups"
OUTPUT_XLSX = ROOT / "excel_output" / "multi_size_first_test" / "eval" / "percentile_5_groups.xlsx"

BATCH = 4
SHEETS = {"mAP50": "0.0000", "mAP50-95": "0.0000", "precision": "0.0000", "recall": "0.0000", "scale": "0.00"}  # sheet -> number format


def load_design():
    """Returns:
        (list[int]): imgsz values, ascending.
        (list[float]): group bin edges (len = n_imgsz + 1).
    """
    design = pd.read_excel(RES_XLSX, sheet_name="design")
    return design["imgsz"].tolist(), [design["lower"].iloc[0]] + design["upper"].tolist()


def load_val_resolutions():
    """Returns:
        (pd.DataFrame): val images only, columns filestem, max_dim.
    """
    df = pd.read_excel(RES_XLSX, sheet_name="resolution")
    return df.loc[df["split"] == "val", ["filestem", "max_dim"]].reset_index(drop=True)


def assign_groups(df, imgsz_list, edges):
    """Add a 'group' column = the imgsz whose design bin contains max_dim.

    Returns:
        (pd.DataFrame): input + 'group' column (int imgsz).
    """
    df["group"] = pd.cut(df["max_dim"], bins=edges, labels=imgsz_list, include_lowest=True).astype(int)
    print(df.groupby("group")["max_dim"].agg(["count", "min", "median", "max"]).to_string())
    return df


def write_group_yamls(df, groups):
    """Write one image-list .txt + one data .yaml per group.

    Returns:
        (dict): {group: yaml path}.
    """
    GROUP_DIR.mkdir(parents=True, exist_ok=True)
    yamls = {}
    for g in groups:
        stems = df.loc[df["group"] == g, "filestem"]
        txt_path = GROUP_DIR / f"val_{g}.txt"
        txt_path.write_text("".join(f"{SPLIT_DIR / 'val' / 'images' / f'{s}.png'}\n" for s in stems))

        yaml_path = GROUP_DIR / f"val_{g}.yaml"
        yaml_path.write_text(
            f"train: {SPLIT_DIR / 'train' / 'images'}\n"  # required key, not used by val
            f"val: {txt_path}\n"
            f"nc: 1\n"
            f"names: ['fracture']\n"
        )
        yamls[g] = yaml_path
    return yamls


def evaluate(group_yamls, group_stats, imgsz_list):
    """Run model.val() for every (imgsz, group) pair with best.pt.

    Args:
        group_stats (pd.DataFrame): index = group, columns count, median (of max_dim).

    Returns:
        (pd.DataFrame): one row per (group, imgsz) with scale, P, R, mAP50, mAP50-95.
    """
    rows = []
    for imgsz in imgsz_list:
        model = YOLO(str(SWEEP_DIR / f"imgsz{imgsz}" / "weights" / "best.pt"))
        for g in imgsz_list:
            print(f"\n=== imgsz {imgsz} | group {g} ===")
            m = model.val(
                data=str(group_yamls[g]), imgsz=imgsz, batch=BATCH, split="val",
                plots=False, project=str(EVAL_DIR), name=f"imgsz{imgsz}_group{g}", exist_ok=True,
            )
            rows.append({
                "group": g,
                "n_images": int(group_stats.loc[g, "count"]),
                "imgsz": imgsz,
                "scale": round(imgsz / group_stats.loc[g, "median"], 2),  # <1 shrunk, >1 enlarged
                "precision": round(m.box.mp, 4),
                "recall": round(m.box.mr, 4),
                "mAP50": round(m.box.map50, 4),
                "mAP50-95": round(m.box.map, 4),
            })
    return pd.DataFrame(rows).sort_values(["group", "imgsz"]).reset_index(drop=True)


def export_excel(results, edges):
    """Write one sheet per metric: rows = image group, columns = model imgsz, + a raw sheet.

    Idea:
        Diagonal (group == imgsz) = matched size; left of it = images shrunk,
        right of it = images enlarged. Bold = best value in the row.

    Args:
        results (pd.DataFrame): output of evaluate().
        edges (list[float]): group bin edges from load_design().

    Returns:
        (pathlib.Path): path of the saved .xlsx.
    """
    font = "Times New Roman"
    side = Side(style="thin", color="BFBFBF")
    border = Border(top=side, bottom=side, left=side, right=side)
    fill = {k: PatternFill("solid", start_color=v) for k, v in
            {"header": "D9E1F2", "match": "BDD7EE", "shrink": "FCE4D6", "enlarge": "FFF2CC"}.items()}

    groups = sorted(results["group"].unique())
    imgszs = sorted(results["imgsz"].unique())
    ranges = [f"{int(edges[i]) + (i > 0)} – {int(edges[i + 1])}" for i in range(len(groups))]  # integer px range
    n_images = results.drop_duplicates("group").set_index("group")["n_images"]

    wb = Workbook()
    wb.remove(wb.active)
    for metric, num_fmt in SHEETS.items():
        ws = wb.create_sheet(metric)
        ws["A1"] = f"{metric} (val set, best.pt). Rows = image group named after its matching imgsz; columns = model imgsz."
        ws["A2"] = "Blue = matched (scale ~ 1). Orange = images shrunk. Yellow = images enlarged. Bold = best in the row."
        for ref in ("A1", "A2"):
            ws[ref].font = Font(name=font, italic=True, color="595959")

        header = ["group (imgsz)", "max_dim (px)"] + [int(s) for s in imgszs] + ["n_images"]
        for j, h in enumerate(header, 1):
            c = ws.cell(4, j, h)
            c.font, c.fill, c.border = Font(name=font, bold=True), fill["header"], border
            c.alignment = Alignment(horizontal="center")

        table = results.pivot(index="group", columns="imgsz", values=metric)
        for i, g in enumerate(groups):
            r = 5 + i
            best = table.loc[g].max()
            row = [int(g), ranges[i]] + [float(table.loc[g, s]) for s in imgszs] + [int(n_images[g])]
            for j, v in enumerate(row, 1):
                c = ws.cell(r, j, v)
                c.border = border
                c.font = Font(name=font, bold=(j == 1), color="808080" if j in (2, len(row)) else "000000")
                if 3 <= j <= len(row) - 1:  # metric cells
                    s = imgszs[j - 3]
                    c.number_format = num_fmt
                    c.alignment = Alignment(horizontal="center")
                    c.fill = fill["match"] if s == g else fill["shrink"] if s < g else fill["enlarge"]
                    if metric != "scale" and v == best:
                        c.font = Font(name=font, bold=True)

        ws.column_dimensions["A"].width = 16
        ws.column_dimensions["B"].width = 14
        for j in range(3, len(header) + 1):
            ws.column_dimensions[get_column_letter(j)].width = 11
        ws.freeze_panes = "C5"

    raw = wb.create_sheet("raw")
    raw.append(list(results.columns))
    for rec in results.itertuples(index=False):
        raw.append([v.item() if hasattr(v, "item") else v for v in rec])

    OUTPUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_XLSX)
    return OUTPUT_XLSX


if __name__ == "__main__":
    imgsz_list, edges = load_design()
    val_df = assign_groups(load_val_resolutions(), imgsz_list, edges)
    group_yamls = write_group_yamls(val_df, imgsz_list)
    group_stats = val_df.groupby("group")["max_dim"].agg(["count", "median"])

    results = evaluate(group_yamls, group_stats, imgsz_list)
    print("\n" + results.to_string(index=False))
    print(f"\nSaved: {export_excel(results, edges)}")