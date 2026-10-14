#!/usr/bin/env python3
"""Convert every results.csv under runs/ into an Excel file, and summarize
the best value of every "metrics/..." column per run + the epoch it happened
at, into a performance_<group>_summary.xlsx — with a "Notes" column listing
the training config (model/epochs/imgsz/batch/device) read from each run's
args.yaml.

Idea:
    Runs are grouped by the folder they sit in under runs/, each group gets
    its own sub-folder in excel_output/ (see GROUP_OUTPUT_NAMES) with its own
    per-run .xlsx files + its own performance_<group>_summary.xlsx -> bold
    "best" values only compare runs inside the same experiment.

If 'model' in args.yaml is a checkpoint PATH instead of a bare pretrained
name, the script loads the checkpoint and counts its parameters to
auto-detect the real size (n/s/m/l/x).

Assumptions (flag if wrong):
  - The epoch column in results.csv is literally named "epoch".
  - Every column whose header starts with "metrics/" is higher-is-better.
  - args.yaml sits in the same folder as results.csv (Ultralytics default).
"""

import csv
import re
from pathlib import Path

import torch
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

RUNS_DIR = Path.home() / "VINH_TRUONG" / "yolo26" / "runs"
OUTPUT_DIR = Path.home() / "VINH_TRUONG" / "yolo26" / "excel_output"

# optional rename: folder under runs/ -> sub-folder name in excel_output/
# empty = use the runs/ folder name as-is
GROUP_OUTPUT_NAMES = {}

METRIC_PREFIX = "metrics/"  # Ultralytics' own convention for performance columns
NOTE_KEYS = ["model", "epochs", "imgsz", "batch", "device"]  # edit this list to change Notes content

# Official param counts (fused, Ultralytics docs), used to auto-detect model size
MODEL_SIZE_REFERENCE = {"n": 2.4e6, "s": 9.5e6, "m": 20.4e6, "l": 24.8e6, "x": 55.7e6}


def resolve_model_name(model_value: str) -> str:
    """If model_value is a checkpoint path (contains '/' or '\\'), open it,
    count parameters, and match against MODEL_SIZE_REFERENCE to recover the
    real size. If it's already a bare name like 'yolo26m.pt', return as-is."""
    if "/" not in model_value and "\\" not in model_value:
        return model_value

    try:
        ckpt = torch.load(model_value, map_location="cpu", weights_only=False)
        n_params = sum(p.numel() for p in ckpt["model"].parameters())
        closest = min(MODEL_SIZE_REFERENCE, key=lambda k: abs(MODEL_SIZE_REFERENCE[k] - n_params))
        return f"yolo26{closest}.pt"
    except Exception as e:
        print(f"  WARNING: could not resolve model size from '{model_value}' ({e})")
        return model_value  # fall back to the raw path if anything goes wrong


def clean_metric_name(header: str) -> str:
    """Strip the 'metrics/' prefix and a trailing '(B)'/'(M)' task suffix."""
    name = header[len(METRIC_PREFIX):] if header.startswith(METRIC_PREFIX) else header
    return re.sub(r"\(\w+\)$", "", name).strip()


def natural_sort_key(run_name: str):
    """Compare every number inside the name by value, not by text
    (imgsz576 < imgsz1024, baseline-2 < baseline-10)."""
    return [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", run_name)]


def read_and_clean_csv(csv_path: Path) -> list:
    """Read results.csv, converting numeric cells to float. Row 0 stays strings (header)."""
    with open(csv_path, newline="") as f:
        raw_rows = list(csv.reader(f))
    cleaned_rows = []
    for row in raw_rows:
        cleaned = []
        for cell in row:
            cell = cell.strip()
            try:
                cleaned.append(float(cell))
            except ValueError:
                cleaned.append(cell)
        cleaned_rows.append(cleaned)
    return cleaned_rows


def read_args_yaml(args_path: Path) -> dict:
    """Read args.yaml -> {key: value}. Tries pyyaml first, falls back to a
    minimal manual 'key: value' line parser if pyyaml isn't installed."""
    if not args_path.exists():
        return {}

    try:
        import yaml
        with open(args_path) as f:
            data = yaml.safe_load(f)
        return data if isinstance(data, dict) else {}
    except ImportError:
        pass  # fall through to manual parsing below

    result = {}
    with open(args_path) as f:
        for line in f:
            line = line.rstrip("\n")
            if not line or line.lstrip().startswith("#") or ":" not in line:
                continue
            key, _, value = line.partition(":")
            key = key.strip()
            value = value.strip()
            if key and not key[0].isspace():  # only top-level keys, skip nested blocks
                result[key] = value
    return result


def build_notes(args: dict) -> str:
    parts = []
    for key in NOTE_KEYS:
        if key not in args or args[key] in (None, ""):
            continue
        value = args[key]
        if key == "model":
            value = resolve_model_name(str(value))
        parts.append(f"{key}={value}")
    return ", ".join(parts)


def write_run_xlsx(rows: list, xlsx_path: Path) -> None:
    """Write the cleaned csv rows to a styled .xlsx."""
    wb = Workbook()
    ws = wb.active
    ws.title = "results"

    for row in rows:
        ws.append(row)

    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"

    for col_idx in range(1, ws.max_column + 1):
        max_len = max(
            (len(str(ws.cell(row=r, column=col_idx).value)) for r in range(1, ws.max_row + 1)),
            default=8,
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 24)

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)


def summarize_run(rows: list) -> dict:
    """Return {metric_name: (best_value, best_epoch)} from cleaned csv rows."""
    if len(rows) < 2:
        return {}
    header = [str(h).strip() if not isinstance(h, float) else str(h) for h in rows[0]]
    data_rows = rows[1:]

    try:
        epoch_col = header.index("epoch")
    except ValueError:
        print("  WARNING: no 'epoch' column, skipping summary for this run")
        return {}

    metric_cols = [(i, h) for i, h in enumerate(header) if h.startswith(METRIC_PREFIX)]
    if not metric_cols:
        print(f"  WARNING: no '{METRIC_PREFIX}*' columns, skipping summary for this run")
        return {}

    result = {}
    for col_idx, raw_name in metric_cols:
        best_value = None
        best_epoch = None
        for row in data_rows:
            value = row[col_idx]
            epoch = row[epoch_col]
            if not isinstance(value, float) or not isinstance(epoch, float):
                continue
            if best_value is None or value > best_value:
                best_value = value
                best_epoch = epoch
        if best_value is not None:
            result[clean_metric_name(raw_name)] = (best_value, best_epoch)
    return result


def write_summary_xlsx(run_summaries: dict, run_notes: dict, all_metrics: list, summary_path: Path) -> None:
    sorted_names = sorted(run_summaries, key=natural_sort_key)

    wb = Workbook()
    ws = wb.active
    ws.title = "best_epoch_summary"

    ws.append(["Name", "Notes"] + all_metrics)
    for run_name in sorted_names:
        summary = run_summaries[run_name]
        row = [run_name, run_notes.get(run_name, "")]
        for metric in all_metrics:
            if metric in summary:
                value, epoch = summary[metric]
                row.append(f"{value:.4g}({int(epoch)})")
            else:
                row.append("")
        ws.append(row)

    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"

    # Bold the best (highest) value in each metric column (col 1 = Name, col 2 = Notes).
    for col_idx, metric in enumerate(all_metrics, start=3):
        best_row = None
        best_value = None
        for row_idx, run_name in enumerate(sorted_names, start=2):
            summary = run_summaries[run_name]
            if metric not in summary:
                continue
            value, _epoch = summary[metric]
            if best_value is None or value > best_value:
                best_value = value
                best_row = row_idx
        if best_row is not None:
            ws.cell(row=best_row, column=col_idx).font = Font(bold=True)

    for col_idx in range(1, ws.max_column + 1):
        max_len = max(
            (len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(1, ws.max_row + 1)),
            default=8,
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 40)

    wb.save(summary_path)


def main():
    if not RUNS_DIR.exists():
        print(f"Directory not found: {RUNS_DIR}")
        return

    csv_files = sorted(RUNS_DIR.glob("**/results.csv"))
    if not csv_files:
        print(f"No results.csv files found under {RUNS_DIR}")
        return

    print(f"Found {len(csv_files)} results.csv file(s):\n")
    groups = {}  # out_dir -> {"summaries", "notes", "metrics"}

    for csv_path in csv_files:
        folder_name = csv_path.parent.name
        group = str(csv_path.parent.parent.relative_to(RUNS_DIR))
        out_dir = OUTPUT_DIR / GROUP_OUTPUT_NAMES.get(group, group)
        xlsx_path = out_dir / f"{folder_name}_results.xlsx"

        rows = read_and_clean_csv(csv_path)
        if not rows:
            print(f"  (empty, skipping) {csv_path}")
            continue

        write_run_xlsx(rows, xlsx_path)

        g = groups.setdefault(out_dir, {"summaries": {}, "notes": {}, "metrics": []})
        summary = summarize_run(rows)
        g["summaries"][folder_name] = summary
        for metric in summary:
            if metric not in g["metrics"]:
                g["metrics"].append(metric)

        args = read_args_yaml(csv_path.parent / "args.yaml")
        g["notes"][folder_name] = build_notes(args)
        if not args:
            print(f"  WARNING: no args.yaml found for {folder_name}, Notes will be empty")

        print(f"  {csv_path.relative_to(RUNS_DIR)}  ->  {xlsx_path.relative_to(OUTPUT_DIR.parent)}  ({len(summary)} metric(s))")

    for out_dir, g in groups.items():
        summary_path = out_dir / f"performance_{out_dir.name}_summary.xlsx"
        write_summary_xlsx(g["summaries"], g["notes"], g["metrics"], summary_path)
        print(f"  summary  ->  {summary_path.relative_to(OUTPUT_DIR.parent)}")

    print(f"\nDone. {len(csv_files)} per-run Excel file(s) + {len(groups)} performance_<group>_summary.xlsx saved under {OUTPUT_DIR}")


if __name__ == "__main__":
    main()