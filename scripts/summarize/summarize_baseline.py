"""Collect the baseline evaluation results into one Excel table.

Reads runs/eval/baseline/<run>/metrics.yaml (written by eval/eval_baseline.py)
and writes all runs as one table to excel_output/eval_baseline.xlsx.

Usage:
    python summarize_baseline.py
"""

from __future__ import annotations

import logging
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ===== SETTINGS (edit here) =====
ROOT = Path.home() / "VINH_TRUONG" / "yolo26"
EVAL_DIR = ROOT / "runs" / "eval" / "baseline"
OUTPUT_XLSX = ROOT / "excel_output" / "eval_baseline.xlsx"

LOGGER = logging.getLogger("summarize_baseline")
COLUMNS = ["run", "model", "imgsz", "train_batch", "seed", "epochs", "split", "val_batch", "n_images",
           "precision", "recall", "mAP50", "mAP50-95"]
METRICS = {"precision", "recall", "mAP50", "mAP50-95"}
SIZE_ORDER = "nsmlx"  # YOLO26 sizes from smallest to largest


def read_results(eval_dir: Path) -> list[dict]:
    """Read every metrics.yaml under the evaluation folder.

    Args:
        eval_dir: Folder with one sub-folder per evaluated run.

    Returns:
        One dict per run, sorted by model size (n, s, m, l, x).
    """
    rows = [yaml.safe_load(f.read_text()) for f in eval_dir.glob("*/metrics.yaml")]
    return sorted(rows, key=size_rank)


def size_rank(row: dict) -> tuple[int, str]:
    """Sort key that puts models in the order n, s, m, l, x.

    Args:
        row: One run, with a "model" value such as "yolo26n".

    Returns:
        Position of the size letter (unknown sizes last), then the run name.
    """
    letter = str(row["model"])[-1]
    rank = SIZE_ORDER.index(letter) if letter in SIZE_ORDER else len(SIZE_ORDER)
    return rank, str(row["run"])


def write_excel(rows: list[dict], path: Path) -> None:
    """Write all runs as one table with a filter on the header row.

    Args:
        rows: One dict per run, keys as in COLUMNS.
        path: Output .xlsx file.
    """
    font = "Times New Roman"
    thin = Side(style="thin", color="BFBFBF")
    border = Border(top=thin, bottom=thin, left=thin, right=thin)
    center = Alignment(horizontal="center")

    wb = Workbook()
    ws = wb.active
    ws.title = "eval_baseline"
    first = rows[0]
    ws["A1"] = (f"Baseline runs, best.pt on the {first['split']} set "
                f"(model.val, batch={first['val_batch']}). Source: runs/eval/baseline/<run>/metrics.yaml")
    ws["A1"].font = Font(name=font, italic=True, color="595959")

    for j, name in enumerate(COLUMNS, start=1):
        c = ws.cell(3, j, name)
        c.font = Font(name=font, bold=True)
        c.fill = PatternFill("solid", fgColor="D9E1F2")
        c.alignment = center
        c.border = border
    for i, row in enumerate(rows, start=4):
        for j, name in enumerate(COLUMNS, start=1):
            value = row.get(name)
            if name in METRICS:
                value = round(value, 3)  # full values stay in metrics.yaml
            c = ws.cell(i, j, value)
            c.font = Font(name=font)
            c.alignment = center
            c.border = border
            if name in METRICS:
                c.number_format = "0.000"  # keeps a trailing zero, e.g. 0.920

    last_col = get_column_letter(len(COLUMNS))
    ws.auto_filter.ref = f"A3:{last_col}{3 + len(rows)}"
    for j in range(1, len(COLUMNS) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 14
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> None:
    """Build excel_output/eval_baseline.xlsx from the evaluation results."""
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    rows = read_results(EVAL_DIR)
    if not rows:
        LOGGER.error("No metrics.yaml found in %s; run eval/eval_baseline.py first", EVAL_DIR)
        return
    write_excel(rows, OUTPUT_XLSX)
    LOGGER.info("Saved %s (%d runs)", OUTPUT_XLSX, len(rows))


if __name__ == "__main__":
    main()